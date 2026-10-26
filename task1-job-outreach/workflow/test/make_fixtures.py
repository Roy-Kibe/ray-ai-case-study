"""Export the Settings and Prompts tabs of the Sheet template as JSON fixtures for run_local.js.

Usage: python make_fixtures.py <output_dir>
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

TEMPLATE = Path(__file__).parents[2] / "sheet-template" / "RAY_Lead_Engine_Template.xlsx"
out = Path(sys.argv[1])
out.mkdir(parents=True, exist_ok=True)

wb = load_workbook(TEMPLATE)
for tab, filename in [("Settings", "settings.json"), ("Prompts", "prompts.json")]:
    ws = wb[tab]
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    (out / filename).write_text(json.dumps(rows, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{filename}: {len(rows)} rows")
